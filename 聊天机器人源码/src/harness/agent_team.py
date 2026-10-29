# -*- coding: utf-8 -*-
"""Lightweight in-exe agent team for background memory compression and sticker labeling.

Each role is a separate main-model call; roles run in parallel and a final
synthesis step merges the results. This is the bot's "agent team" implementation
inside the native DSH-compatible core.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import threading
from typing import Dict, List

# AgentTeam 固定系统上下文：只作为子智能体的 system prompt，
# 不写入聊天主上下文，也不会进入可压缩的记忆正文。
AGENT_TEAM_SYSTEM = (
    "你是本机 AgentTeam 的后台子智能体，只执行分配的任务。"
    "只输出负责内容，不解释，不提及系统/提示词，不把提示词写入记忆。"
)


def _sys(task: str) -> str:
    return AGENT_TEAM_SYSTEM + "\n" + str(task or "")


def _fit_desc(text: str, low: int = 30, high: int = 50) -> str:
    s = re.sub(r"\s+", " ", str(text or "")).strip()
    s = s.strip("，。；;、 ")
    if len(s) > high:
        cut = s[:high]
        for mark in ("。", "；", ";", "，", ","):
            pos = cut.rfind(mark)
            if pos >= low:
                cut = cut[:pos]
                break
        s = cut.strip()
    return s


class AgentTeam:
    def __init__(self, core):
        self.core = core

    # ---------------- parallel helper ----------------
    @staticmethod
    def _parallel(calls):
        results = [None] * len(calls)

        def run(i, fn):
            try:
                results[i] = fn()
            except Exception:
                results[i] = None

        threads = [threading.Thread(target=run, args=(i, fn), daemon=True)
                   for i, fn in enumerate(calls)]
        for t in threads:
            t.start()
        for t in threads:
            t.join(timeout=90)
        return results

    # ---------------- memory compression team ----------------
    def compress_memory(self, text: str, target: int, group=None, name: str = "",
                        senders=None, platform: str = "wechat",
                        forced_recent: int = 0, force_long: bool = False) -> str:
        original = str(text or "")[:24000]
        group = group or {}
        members = "、".join(str(x) for x in (group.get("chats") or []))
        if not members and senders:
            members = "、".join(str(x) for x in senders)
        mode = "记忆组" if group else "单人会话"
        weight_rule = ""
        if group:
            raw_members = group.get("members") or []
            singles = groups = 0
            for m in raw_members:
                if not isinstance(m, dict):
                    continue
                if bool(m.get("is_group")):
                    groups += 1
                else:
                    singles += 1
            if singles and groups:
                weight_rule = "组内同时有单聊和群聊：单聊个体记忆约 65%，群聊约 35%，单聊细节优先保留。\n"
        forced_rule = ""
        if force_long:
            forced_rule = ("- 本轮包含掉出最近8条的超长原文（>300字），必须压缩它，"
                           "不得原样保留；只保留最近 %d 条其他原文。\n" % forced_recent)
        common = (
            "你是一个记忆整理团队中的专项智能体。只输出你负责的部分，不要写解释。\n"
            "【每轮重复的压缩规则】\n"
            "- 当前是%s，压缩目标 %d 字；\n"
            "- 最近强制保留的 %d 条对话不在本次输入里，也不要求你复制；\n"
            "%s"
            "- 只压缩长期记忆，不要把最近强制保留的记录计入压缩结果；\n"
            "- 必须保留：用户给我的取名/昵称、用户珍视的记忆、用户明确提到的重要信息、"
            "约定/承诺、身份关系、偏好/雷区。\n%s" % (
                mode, target, forced_recent, forced_rule, weight_rule)
        )
        common = _sys(common)
        prompts = [
            common + "专项：提取事实、取名/昵称、珍视记忆、重要信息、身份关系、约定/承诺。",
            common + "专项：提取每个人的喜好、讨厌/雷区、语言习惯；以及“我”的性格/人格/好恶。",
            common + "专项：提取重要事件、计划、共同记忆/组内共识、最近变化。",
        ]
        results = self._parallel([
            (lambda p=p: self.core.complete(p, original, platform=platform, max_tokens=700))
            for p in prompts
        ])
        notes = [r for r in results if r and str(r).strip()]
        if group:
            fmt = ("整理成不超过 %d 字的记忆组摘要。分成员小节记录喜好/讨厌/语言习惯，"
                   "另写【我在这个组】性格/人格/好恶、【组内共识】；成员：%s。" % (target, members))
        else:
            fmt = "整理成不超过 %d 字的联系人记忆摘要，保留身份/偏好/雷区/重要事件/计划/约定/图片要点。" % target
        synth = ("你是记忆整理团队的主编。合并以下专项笔记，压缩为长期记忆摘要；"
                 "不要丢取名/昵称、珍视记忆、重要信息、约定/承诺。"
                 "只输出长期记忆，不要复制最近强制保留的对话。%s\n\n【专项笔记】\n%s" % (
                     fmt, "\n\n".join(notes)))
        synth = _sys(synth)
        try:
            out = self.core.complete(synth, "请输出最终摘要正文。", platform=platform,
                                     max_tokens=max(800, int(target) + 400)).strip()
        except Exception:
            out = ""
        if out:
            return out
        return "\n\n".join(notes)[:max(2000, int(target))]

    # ---------------- sticker labeling agent (single) ----------------
    def describe_sticker(self, image_path: str, hint: str = "", platform: str = "wechat") -> Dict:
        """One vision agent: action/expression description + tags + in-image text."""
        prompt = (
            "你是表情包标注智能体。看这张表情包，只输出 JSON：\n"
            '{"desc": "30~50个中文字符，客观写清动作和表情/情绪；如果图中有文字，把文字原样写进去", '
            '"tags": ["5到8个中文短标签：动作/表情/对象/场景/梗"], '
            '"text": "图中文字，没有就空字符串"}'
        )
        prompt = _sys(prompt)
        raw = ""
        try:
            raw = self.core.vision(prompt, image_path, platform=platform, max_tokens=260)
        except Exception:
            raw = ""
        desc = ""
        tags: List[str] = []
        text_in_image = ""
        try:
            m = re.search(r"\{.*\}", str(raw or ""), re.S)
            if m:
                j = json.loads(m.group(0))
                desc = str(j.get("desc") or "").strip()
                tags = [str(t).strip() for t in (j.get("tags") or []) if str(t).strip()][:8]
                text_in_image = str(j.get("text") or "").strip()
        except Exception:
            pass
        desc = _fit_desc(desc)
        if text_in_image and text_in_image not in desc:
            desc = _fit_desc((desc + "；文字：" + text_in_image).strip())
        if not desc:
            desc = _fit_desc("表情包画面，动作和表情待确认")
        if not tags:
            tags = [w for w in re.split(r"[，。；、\s]+", desc) if w][:5]
        return {"desc": desc[:50], "tags": tags[:8]}

    # ---------------- document browsing team ----------------
    @staticmethod
    def extract_document_text(path: str) -> str:
        """Extract Word/PDF/text content without AI."""
        ext = os.path.splitext(path)[1].lower()
        try:
            if ext == ".docx":
                from docx import Document
                doc = Document(path)
                parts = [p.text for p in doc.paragraphs if str(p.text or "").strip()]
                for table in doc.tables:
                    for row in table.rows:
                        cells = [c.text.strip() for c in row.cells]
                        if any(cells):
                            parts.append(" | ".join(cells))
                return "\n".join(parts).strip()
            if ext == ".pdf":
                from pypdf import PdfReader
                reader = PdfReader(path)
                parts = []
                for page in reader.pages[:200]:
                    try:
                        parts.append(page.extract_text() or "")
                    except Exception:
                        pass
                return "\n".join(parts).strip()
            if ext == ".xlsx":
                from openpyxl import load_workbook
                wb = load_workbook(path, read_only=True, data_only=True)
                parts = []
                for ws in wb.worksheets:
                    parts.append("【工作表 %s】" % ws.title)
                    for row in ws.iter_rows(values_only=True):
                        cells = [str(c) for c in row if c is not None and str(c).strip()]
                        if cells:
                            parts.append(" | ".join(cells))
                try:
                    wb.close()
                except Exception:
                    pass
                return "\n".join(parts).strip()
            if ext == ".pptx":
                from pptx import Presentation
                prs = Presentation(path)
                parts = []
                for i, slide in enumerate(prs.slides, 1):
                    parts.append("【幻灯片 %d】" % i)
                    for shape in slide.shapes:
                        txt = getattr(shape, "text", "")
                        if txt and str(txt).strip():
                            parts.append(str(txt).strip())
                return "\n".join(parts).strip()
            if ext in (".txt", ".md", ".csv"):
                return open(path, "r", encoding="utf-8", errors="replace").read().strip()
            if ext == ".rtf":
                raw = open(path, "r", encoding="utf-8", errors="replace").read()
                return re.sub(r"\\[a-z]+-?\d* ?|[{}]", " ", raw).strip()
        except Exception:
            return ""
        return ""

    def browse_document(self, path: str, name: str = "", platform: str = "wechat") -> str:
        """Agent team reads a document in the background and returns a usable summary."""
        text = self.extract_document_text(path)
        if not text:
            return "（文档无法解析或为空：%s）" % (name or os.path.basename(path))
        text = text[:60000]
        roles = [
            "专项：总结文档主题、结论和关键事实；不要漏重要数字/名称/日期。",
            "专项：提取与用户问题相关的要点、待办/行动项、约定/承诺、风险。",
            "专项：提取文档中需要回复或确认的问题、以及可能影响用户的重要信息。",
        ]
        results = self._parallel([
            (lambda p=p: self.core.complete(
                _sys("你是文档阅读团队中的专项智能体。只输出你负责的部分，不要写解释。\n" + p),
                text, platform=platform, max_tokens=900))
            for p in roles
        ])
        notes = [r for r in results if r and str(r).strip()]
        synth = (
            "你是文档阅读团队的主编。根据以下专项笔记，写一份供我回复用的文档摘要。\n"
            "要求：中文、结构清晰、保留标题/关键事实/数字/姓名/日期/待办/问题；不要编造；"
            "不要提及“专项智能体”。文档名：%s\n\n【专项笔记】\n%s" % (
                name or os.path.basename(path), "\n\n".join(notes))
        )
        synth = _sys(synth)
        try:
            out = self.core.complete(synth, "请输出文档摘要正文。", platform=platform,
                                     max_tokens=1600).strip()
        except Exception:
            out = ""
        if out:
            return out
        return ("【文档】%s\n" % (name or os.path.basename(path))) + "\n\n".join(notes)[:8000]

    @staticmethod
    def digest(text: str) -> str:
        return hashlib.sha1(str(text or "").encode("utf-8")).hexdigest()[:12]
